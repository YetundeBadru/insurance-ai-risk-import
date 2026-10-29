"""Risk Import Validator: validate uploaded risk schedules before they reach a policy.

Design rule: the LLM may help MAP COLUMN HEADERS only (it never sees cell data) and
may never decide whether a row is valid. Validity comes from deterministic rules.
"""
import csv
import io
import json
import re
import time
from datetime import date

from . import audit
from .jsonutil import parse_json_object
from .master import CLIENTS, POLICIES
from .metrics import metrics as default_metrics
from .risk_schema import ALIASES, CANONICAL, MAP_PROMPT, normalize_header

_NG = re.compile(r"^[A-Z]{3}\d{3}[A-Z]{2}$")
_KE = re.compile(r"^K[A-Z]{2}\d{3}[A-Z]$")


# ---------- reading files ----------
def read_table(filename: str, raw: bytes):
    """Return (headers, rows) where rows is a list of lists of strings. Blank rows dropped."""
    name = filename.lower()
    if name.endswith(".csv"):
        data = list(csv.reader(io.StringIO(raw.decode("utf-8-sig"))))
    elif name.endswith(".xlsx"):
        from openpyxl import load_workbook
        ws = load_workbook(io.BytesIO(raw), read_only=True, data_only=True).worksheets[0]
        data = [["" if c is None else str(c) for c in row] for row in ws.iter_rows(values_only=True)]
    else:
        raise ValueError("unsupported_file_type (supported: .csv, .xlsx)")
    data = [r for r in data if any(str(c).strip() for c in r)]
    if not data:
        raise ValueError("empty_file")
    return [str(h).strip() for h in data[0]], data[1:]


# ---------- column mapping ----------
def map_columns(headers, llm, metrics=default_metrics):
    mapping, unresolved = {}, []
    for h in headers:
        mapping[h] = ALIASES.get(normalize_header(h))
        if mapping[h] is None:
            unresolved.append(h)
    if unresolved:  # LLM fallback sees ONLY header names
        t0 = time.perf_counter()
        try:
            resp = llm.complete(MAP_PROMPT, json.dumps(unresolved))
            metrics.record_call(getattr(llm, "last_call_seconds", None) or (time.perf_counter() - t0), True,
                                resp.input_tokens, resp.output_tokens)
            proposed = parse_json_object(resp.text)
            used = {v for v in mapping.values() if v}
            for h in unresolved:
                v = proposed.get(h) if isinstance(proposed, dict) else None
                if v in CANONICAL and v not in used:  # validate: known field, no double-mapping
                    mapping[h] = v
                    used.add(v)
        except Exception:
            metrics.record_call(time.perf_counter() - t0, False)
    return mapping


# ---------- deterministic row rules ----------
def _money(s):
    s = re.sub(r"(?i)\b(ngn|kes|ksh|usd)\b|[₦,\s]", "", str(s))
    return float(s) if re.fullmatch(r"\d+(\.\d+)?", s) else None


def validate_row(tenant_id, row, seen):
    r = []
    for f in CANONICAL:
        if not str(row.get(f, "")).strip():
            r.append(f"MISSING:{f}")
    plate = re.sub(r"[\s\-]", "", str(row.get("registration_no", "")).upper())
    if plate and not (_NG.match(plate) or _KE.match(plate)):
        r.append("BAD_PLATE")
    yr = str(row.get("year", "")).strip()
    if yr and not (re.fullmatch(r"\d{4}", yr) and 1980 <= int(yr) <= date.today().year + 1):
        r.append("BAD_YEAR")
    amount = None
    if str(row.get("sum_insured", "")).strip():
        amount = _money(row["sum_insured"])
        if amount is None or amount <= 0:
            r.append("BAD_SUM_INSURED")
            amount = None
    pol = str(row.get("policy_number", "")).strip().upper()
    owner = str(row.get("owner_id", "")).strip().upper()
    if pol and plate:
        if (pol, plate) in seen:
            r.append("DUPLICATE_IN_FILE")
        seen.add((pol, plate))
    p = POLICIES.get((tenant_id, pol)) if pol else None
    if pol and not p:
        r.append("POLICY_NOT_FOUND")  # identical for "doesn't exist" and "belongs to another insurer"
    if p and p["status"] != "active":
        r.append("POLICY_INACTIVE")
    if owner and (tenant_id, owner) not in CLIENTS:
        r.append("UNKNOWN_CLIENT")
    elif owner and p and owner != p["owner_id"]:
        r.append("OWNER_MISMATCH")
    if p and amount is not None and amount > p["limit"]:
        r.append("SUM_EXCEEDS_LIMIT")
    return r


def _mask(v):
    v = str(v)
    return v[:2] + "*" * (len(v) - 4) + v[-2:] if len(v) > 4 else "****"


def _public(row):
    out = dict(row)
    if out.get("owner_id"):
        out["owner_id"] = _mask(out["owner_id"])  # masked by default
    return out


def import_risks(tenant_id, filename, raw, llm, dry_run=False, metrics=default_metrics):
    try:
        headers, rows = read_table(filename, raw)
    except ValueError as e:
        audit.record(tenant_id, "risk_import", file_sha256=audit.file_hash(raw), status="rejected", reason=str(e))
        return {"status": "rejected", "reason": str(e), "total": 0, "accepted": [], "rejected": [],
                "column_mapping": {}, "not_processed": 0}

    mapping = map_columns(headers, llm, metrics)
    missing = [f for f in CANONICAL if f not in mapping.values()]
    if missing:  # can't validate safely -> human review, nothing imported
        audit.record(tenant_id, "risk_import", file_sha256=audit.file_hash(raw), status="needs_review",
                     reason="unmapped_required:" + ",".join(missing), rows_total=len(rows))
        return {"status": "needs_review", "reason": "unmapped_required:" + ",".join(missing), "total": len(rows),
                "accepted": [], "rejected": [], "column_mapping": mapping, "not_processed": len(rows)}

    idx = {i: mapping[h] for i, h in enumerate(headers) if mapping[h]}
    accepted, rejected, seen = [], [], set()
    for n, cells in enumerate(rows, start=1):
        row = {idx[i]: (cells[i] if i < len(cells) else "") for i in idx}
        reasons = validate_row(tenant_id, row, seen)
        (rejected if reasons else accepted).append(
            {"row": n, **({"reasons": reasons} if reasons else {}), "data": _public(row)})

    status = "imported" if not rejected else ("rejected" if not accepted else "partial")
    audit.record(tenant_id, "risk_import", file_sha256=audit.file_hash(raw), status=status, dry_run=dry_run,
                 rows_total=len(rows), accepted=len(accepted), rejected=len(rejected))
    return {"status": status, "total": len(rows), "accepted": accepted, "rejected": rejected,
            "column_mapping": mapping, "not_processed": 0, "dry_run": dry_run}
